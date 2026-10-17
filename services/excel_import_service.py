import io
import re
import csv
from datetime import date, timedelta
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession

from models.parking_card import ParkingCard, CardStatus
from models.subscriber import Subscriber
from models.subscription import Subscription, SubscriptionStatus
from models.subscription_plan import SubscriptionPlan
from utils.time import cairo_now


class ExcelImportService:
    @staticmethod
    def _normalize_string(val: Any) -> str:
        if val is None:
            return ""
        s = str(val).strip()
        # Handle excel float numbers like 1001.0 -> 1001
        if re.match(r"^\d+\.0$", s):
            s = s.split(".")[0]
        return s

    @staticmethod
    def _match_column_name(header: str) -> Optional[str]:
        h = header.strip().lower()
        h_clean = re.sub(r"[\s_\-\(\)\.]+", "", h)

        # Card code variations
        if any(k in h_clean for k in ["رقمكارت", "كودكارت", "كارت", "بطاقة", "cardcode", "cardnum", "cardno", "card", "rfid", "barcode"]):
            return "card_code"

        # Company variations
        if any(k in h_clean for k in ["شركة", "شركه", "بنك", "مؤسسة", "مؤسسه", "جهة", "جهه", "company", "bank", "corp", "org", "employer"]):
            return "company_name"

        # Plate number variations
        if any(k in h_clean for k in ["رقملوحة", "رقملوحه", "لوحة", "لوحه", "سيارة", "سياره", "plate", "license", "car"]):
            return "plate_number"

        # Full name variations
        if any(k in h_clean for k in ["اسمالموظف", "اسمالمشترك", "اسمالعميل", "الاسم", "موظف", "مشترك", "عميل", "fullname", "name", "employee", "subscriber"]):
            return "full_name"

        # Phone variations
        if any(k in h_clean for k in ["تليفون", "هاتف", "موبايل", "جوال", "phone", "mobile", "tel"]):
            return "phone_number"

        # Notes variations
        if any(k in h_clean for k in ["ملاحظات", "ملاحظة", "notes", "note", "comment"]):
            return "notes"

        return None

    @classmethod
    def parse_file(cls, file_content: bytes, filename: str) -> List[Dict[str, Any]]:
        """Parses an uploaded Excel or CSV file into standardized row dictionaries."""
        filename_lower = filename.lower()
        raw_rows: List[List[Any]] = []

        if filename_lower.endswith(".csv"):
            text_data = None
            for enc in ["utf-8-sig", "utf-8", "cp1256", "latin1"]:
                try:
                    text_data = file_content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if not text_data:
                raise ValueError("فشل قراءة ترميز ملف CSV. يرجى التأكد من حفظ الملف بترميز UTF-8.")

            reader = csv.reader(io.StringIO(text_data))
            for row in reader:
                if any(str(cell).strip() for cell in row):
                    raw_rows.append(row)

        elif filename_lower.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if row and any(cell is not None and str(cell).strip() != "" for cell in row):
                    raw_rows.append(list(row))
        else:
            raise ValueError("صيغة الملف غير مدعومة. يرجى رفع ملف Excel (.xlsx) أو ملف CSV (.csv)")

        if not raw_rows:
            raise ValueError("الملف المرفوع فارغ ولا يحتوي على أي بيانات.")

        # Find header row (usually the first non-empty row, or search for recognizable headers)
        header_row_idx = 0
        col_map: Dict[str, int] = {}

        for r_idx, row in enumerate(raw_rows[:5]):
            detected = {}
            for c_idx, cell in enumerate(row):
                if cell is not None:
                    matched = cls._match_column_name(str(cell))
                    if matched and matched not in detected:
                        detected[matched] = c_idx
            if "card_code" in detected or len(detected) >= 2:
                header_row_idx = r_idx
                col_map = detected
                break

        if not col_map or "card_code" not in col_map:
            # Fallback: if no clear header found, try column positions (0: card, 1: name, 2: company, 3: plate, 4: phone)
            col_map = {
                "card_code": 0,
                "full_name": 1 if len(raw_rows[0]) > 1 else None,
                "company_name": 2 if len(raw_rows[0]) > 2 else None,
                "plate_number": 3 if len(raw_rows[0]) > 3 else None,
                "phone_number": 4 if len(raw_rows[0]) > 4 else None,
            }
            col_map = {k: v for k, v in col_map.items() if v is not None}
            data_rows = raw_rows
        else:
            data_rows = raw_rows[header_row_idx + 1 :]

        parsed: List[Dict[str, Any]] = []
        for i, row in enumerate(data_rows, start=header_row_idx + 2):
            def get_val(key: str) -> str:
                idx = col_map.get(key)
                if idx is not None and idx < len(row):
                    return cls._normalize_string(row[idx])
                return ""

            card_code = get_val("card_code")
            if not card_code:
                continue

            full_name = get_val("full_name")
            company_name = get_val("company_name")
            plate_number = get_val("plate_number")
            phone_number = get_val("phone_number")
            notes = get_val("notes")

            parsed.append(
                {
                    "row_num": i,
                    "card_code": card_code,
                    "full_name": full_name,
                    "company_name": company_name,
                    "plate_number": plate_number,
                    "phone_number": phone_number,
                    "notes": notes,
                }
            )

        if not parsed:
            raise ValueError("لم يتم العثور على أي صفوف صالحة تحتوي على رقم الكارت في الملف.")

        return parsed

    @classmethod
    async def preview(
        cls,
        parsed_rows: List[Dict[str, Any]],
        default_company: str,
        db: AsyncSession,
    ) -> Dict[str, Any]:
        """Analyzes parsed rows against existing cards and subscribers in the database."""
        card_codes = [r["card_code"] for r in parsed_rows]

        # Fetch existing cards
        existing_cards_res = await db.execute(
            select(ParkingCard).where(ParkingCard.card_code.in_(card_codes))
        )
        existing_cards = {c.card_code: c for c in existing_cards_res.scalars().all()}

        # Fetch existing subscribers by card_code
        existing_subs_by_card_res = await db.execute(
            select(Subscriber).where(Subscriber.card_code.in_(card_codes))
        )
        existing_subs_by_card = {s.card_code: s for s in existing_subs_by_card_res.scalars().all()}

        # Also fetch existing subscribers by plate (for those having a plate)
        plates = [r["plate_number"] for r in parsed_rows if r.get("plate_number")]
        existing_subs_by_plate = {}
        if plates:
            existing_subs_by_plate_res = await db.execute(
                select(Subscriber).where(Subscriber.plate_number.in_(plates))
            )
            existing_subs_by_plate = {s.plate_number: s for s in existing_subs_by_plate_res.scalars().all()}

        seen_cards = set()
        items = []
        new_cards_count = 0
        existing_cards_count = 0
        new_subscribers_count = 0
        existing_subscribers_count = 0
        duplicates_in_file_count = 0

        for r in parsed_rows:
            card_code = r["card_code"]
            company = r["company_name"] or default_company
            name = r["full_name"] or (f"موظف {company}" if company else f"مشترك كارت {card_code}")
            plate = r["plate_number"]

            is_dup_in_file = card_code in seen_cards
            seen_cards.add(card_code)

            card_obj = existing_cards.get(card_code)
            sub_obj = existing_subs_by_card.get(card_code) or (existing_subs_by_plate.get(plate) if plate else None)

            status_tag = "NEW"
            status_text = "جديد"
            badge_color = "green"
            warning_msg = None

            if is_dup_in_file:
                status_tag = "DUPLICATE_FILE"
                status_text = "مكرر بالملف"
                badge_color = "red"
                warning_msg = "رقم الكارت مكرر أكثر من مرة داخل هذا الملف"
                duplicates_in_file_count += 1
            elif sub_obj:
                status_tag = "EXISTING_SUBSCRIBER"
                status_text = "مشترك موجود"
                badge_color = "yellow"
                warning_msg = f"المشترك مسجل مسبقاً باسم: {sub_obj.full_name}"
                existing_subscribers_count += 1
            elif card_obj:
                status_tag = "CARD_EXISTS"
                status_text = "الكارت موجود"
                badge_color = "blue"
                warning_msg = f"الكارت مسجل مسبقاً بحالة: {card_obj.status.value}"
                existing_cards_count += 1
                new_subscribers_count += 1
            else:
                new_cards_count += 1
                new_subscribers_count += 1

            items.append(
                {
                    "row_num": r["row_num"],
                    "card_code": card_code,
                    "full_name": name,
                    "company_name": company,
                    "plate_number": plate or "بدون لوحة",
                    "phone_number": r["phone_number"],
                    "notes": r["notes"],
                    "status_tag": status_tag,
                    "status_text": status_text,
                    "badge_color": badge_color,
                    "warning_msg": warning_msg,
                }
            )

        return {
            "total_rows": len(parsed_rows),
            "new_cards_count": new_cards_count,
            "existing_cards_count": existing_cards_count,
            "new_subscribers_count": new_subscribers_count,
            "existing_subscribers_count": existing_subscribers_count,
            "duplicates_in_file_count": duplicates_in_file_count,
            "items": items,
        }

    @classmethod
    async def execute_import(
        cls,
        parsed_rows: List[Dict[str, Any]],
        default_company: str,
        plan_id: Optional[int],
        update_existing: bool,
        current_user_id: int,
        db: AsyncSession,
    ) -> Dict[str, Any]:
        """Performs atomic batch import of cards, subscribers, and optional subscriptions."""
        selected_plan = None
        if plan_id:
            plan_res = await db.execute(select(SubscriptionPlan).where(SubscriptionPlan.id == plan_id))
            selected_plan = plan_res.scalars().first()

        created_cards = 0
        created_subscribers = 0
        updated_subscribers = 0
        created_subscriptions = 0
        skipped_count = 0
        errors = []

        now_cairo = cairo_now()
        start_date = now_cairo.date()

        seen_cards = set()

        for r in parsed_rows:
            card_code = r["card_code"]
            if not card_code or card_code in seen_cards:
                skipped_count += 1
                continue
            seen_cards.add(card_code)

            company = (r.get("company_name") or "").strip() or default_company
            full_name = (r.get("full_name") or "").strip() or (f"موظف {company}" if company else f"مشترك كارت {card_code}")
            plate_number = (r.get("plate_number") or "").strip()
            phone_number = (r.get("phone_number") or "").strip() or None
            notes = (r.get("notes") or "").strip() or None

            # 1. Handle Parking Card
            card_res = await db.execute(select(ParkingCard).where(ParkingCard.card_code == card_code))
            card = card_res.scalars().first()
            if not card:
                card = ParkingCard(
                    card_code=card_code,
                    status=CardStatus.IN_USE if selected_plan else CardStatus.AVAILABLE,
                    created_by=current_user_id,
                )
                db.add(card)
                await db.flush()
                created_cards += 1
            else:
                if selected_plan and card.status != CardStatus.IN_USE:
                    card.status = CardStatus.IN_USE
                    await db.flush()

            # 2. Determine Plate Number (Must be non-empty and unique for SQLite constraint)
            if not plate_number:
                plate_number = f"بدون لوحة-{card_code}"

            # 3. Handle Subscriber
            sub_res = await db.execute(
                select(Subscriber).where(
                    or_(Subscriber.card_code == card_code, Subscriber.plate_number == plate_number)
                )
            )
            subscriber = sub_res.scalars().first()

            if subscriber:
                if update_existing:
                    subscriber.full_name = full_name
                    if company:
                        subscriber.company_name = company
                    if phone_number:
                        subscriber.phone_number = phone_number
                    subscriber.card_code = card_code
                    if notes:
                        subscriber.notes = notes
                    await db.flush()
                    updated_subscribers += 1
                else:
                    skipped_count += 1
            else:
                subscriber = Subscriber(
                    full_name=full_name,
                    plate_number=plate_number,
                    phone_number=phone_number,
                    company_name=company,
                    card_code=card_code,
                    notes=notes,
                )
                db.add(subscriber)
                await db.flush()
                created_subscribers += 1

            # 4. Handle Subscription (if Plan selected)
            if selected_plan and subscriber and card:
                # Check if already has active subscription
                active_sub_res = await db.execute(
                    select(Subscription).where(
                        Subscription.subscriber_id == subscriber.id,
                        Subscription.status == SubscriptionStatus.ACTIVE,
                    )
                )
                existing_active_sub = active_sub_res.scalars().first()

                if not existing_active_sub:
                    end_date = start_date + timedelta(days=selected_plan.duration_days)
                    sub_record = Subscription(
                        subscriber_id=subscriber.id,
                        plan_id=selected_plan.id,
                        card_id=card.id,
                        plate_number=plate_number,
                        start_date=start_date,
                        end_date=end_date,
                        status=SubscriptionStatus.ACTIVE,
                        amount_paid_piastres=selected_plan.price_piastres,
                        plan_price_snapshot=selected_plan.price_piastres,
                        paid_at=now_cairo,
                        collected_by=current_user_id,
                        notes=f"اشتراك جماعي (استيراد إكسيل) - {company}" if company else "اشتراك جماعي (استيراد إكسيل)",
                    )
                    db.add(sub_record)
                    await db.flush()
                    created_subscriptions += 1

        await db.commit()

        return {
            "success": True,
            "created_cards": created_cards,
            "created_subscribers": created_subscribers,
            "updated_subscribers": updated_subscribers,
            "created_subscriptions": created_subscriptions,
            "skipped_count": skipped_count,
            "errors": errors,
        }

    @staticmethod
    def generate_template() -> bytes:
        """Generates a clean, stylized Excel template for corporate subscriber imports."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "المشتركين"
        ws.views.sheetView[0].rightToLeft = True

        headers = [
            "رقم الكارت (إجباري)",
            "اسم الموظف / العميل",
            "اسم الشركة / البنك",
            "رقم اللوحة (اختياري)",
            "رقم الهاتف (اختياري)",
            "ملاحظات",
        ]

        # Style header
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Sample rows
        sample_rows = [
            ["3101", "أحمد محمد محمود", "بنك مصر", "ق ر س 1 2 3", "01012345678", "الفرع الرئيسي"],
            ["3102", "سارة علي إبراهيم", "بنك مصر", "", "01123456789", "قسم الحسابات"],
            ["3103", "محمد حسن خليل", "شركة فودافون", "أ ب ج 4 5 6", "01234567890", "إدارة التسويق"],
            ["3104", "محمود إبراهيم سيد", "شركة فودافون", "", "", "موظف مبيعات"],
        ]

        data_font = Font(name="Segoe UI", size=10)
        data_align = Alignment(horizontal="center", vertical="center")

        for r in sample_rows:
            ws.append(r)
            row_idx = ws.max_row
            for col_num in range(1, len(r) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.alignment = data_align

        # Adjust column widths
        col_widths = [22, 26, 24, 22, 22, 25]
        for idx, width in enumerate(col_widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
